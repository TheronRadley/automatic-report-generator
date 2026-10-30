"""Exporters: PDF/HTML/Excel content, escaping, overwrite and atomicity."""

from __future__ import annotations

from pathlib import Path

import openpyxl
import pytest

from reportgen.config.loader import load_config
from reportgen.errors import ExportError
from reportgen.exporters.base import export_document
from reportgen.reports.builder import build_report, run_report


def _document(minimal_sales_config, tmp_path):
    config, base = load_config(minimal_sales_config)
    return build_report(config, base, tmp_path / "work")


def test_pdf_html_xlsx_generated(minimal_sales_config, tmp_path):
    outputs = run_report(minimal_sales_config, formats=["pdf", "html", "xlsx"])
    by_suffix = {p.suffix: p for p in outputs}
    assert by_suffix[".pdf"].read_bytes().startswith(b"%PDF")
    assert by_suffix[".pdf"].stat().st_size > 5000
    html = by_suffix[".html"].read_text(encoding="utf-8")
    assert "Test Sales Report" in html
    assert "$890.00" in html  # 100+180+150+100+360 total revenue
    assert "data:image/png;base64," in html
    assert "http://" not in html and "https://" not in html  # fully self-contained
    book = openpyxl.load_workbook(by_suffix[".xlsx"])
    assert "Summary" in book.sheetnames
    assert len(book.sheetnames) >= 3


def test_overwrite_refused_by_default(minimal_sales_config, tmp_path):
    document = _document(minimal_sales_config, tmp_path)
    out_dir = tmp_path / "out2"
    out_dir.mkdir()
    export_document(document, "html", out_dir, "r", overwrite=False)
    with pytest.raises(ExportError, match="already exists"):
        export_document(document, "html", out_dir, "r", overwrite=False)
    export_document(document, "html", out_dir, "r", overwrite=True)


def test_html_escapes_user_content(write_config, sales_csv, tmp_path):
    config = write_config(f"""
version: 1
report:
  title: "<script>alert(1)</script>"
  template: general
sources:
  - {{name: d, type: csv, path: "{sales_csv}"}}
exports:
  output_dir: "{tmp_path / "o"}"
  filename: esc
  overwrite: true
""")
    outputs = run_report(config, formats=["html"])
    html = outputs[0].read_text(encoding="utf-8")
    assert "<script>alert(1)</script>" not in html
    assert "&lt;script&gt;" in html


def test_excel_formula_injection_escaped(write_config, tmp_path):
    csv_path = tmp_path / "evil.csv"
    csv_path.write_text("name,value\n=2+2,3\nok,4\n", encoding="utf-8")
    config = write_config(f"""
version: 1
report: {{title: Evil test, template: general}}
sources:
  - {{name: d, type: csv, path: "{csv_path}"}}
sections:
  - {{title: Data, type: table, dataset: d}}
exports:
  output_dir: "{tmp_path / "o"}"
  filename: evil
  overwrite: true
""")
    outputs = run_report(config, formats=["xlsx"])
    book = openpyxl.load_workbook(outputs[0])
    sheet = book["Data"]
    assert sheet.cell(row=2, column=1).value == "'=2+2"


def test_failed_export_leaves_no_partial_file(minimal_sales_config, tmp_path):
    outputs = run_report(minimal_sales_config, formats=["html"])
    good_bytes = outputs[0].read_bytes()
    # A failing run must not touch the previously completed report, and
    # must not leave temporary partial files behind.
    with pytest.raises(ExportError):
        run_report(minimal_sales_config, formats=["docx"])  # type: ignore[list-item]
    assert outputs[0].read_bytes() == good_bytes
    leftovers = [p for p in outputs[0].parent.iterdir() if p.name.startswith(".")]
    assert leftovers == []
    # Earlier formats in the same run are preserved when a later one fails:
    with pytest.raises(ExportError):
        run_report(minimal_sales_config, formats=["pdf", "docx"])  # type: ignore[list-item]
    assert (outputs[0].parent / "test-report.pdf").exists()


def test_empty_dataset_refused(write_config, tmp_path):
    csv_path = tmp_path / "empty.csv"
    csv_path.write_text("a,b\n", encoding="utf-8")
    config = write_config(f"""
version: 1
report: {{title: Empty, template: general}}
sources:
  - {{name: d, type: csv, path: "{csv_path}"}}
exports:
  output_dir: "{tmp_path / "o"}"
  filename: empty
""")
    with pytest.raises(Exception, match="empty"):
        run_report(config, formats=["html"])


def test_unknown_metric_placeholder_errors(write_config, sales_csv, tmp_path):
    config = write_config(f"""
version: 1
report:
  title: Bad placeholder
  template: general
  summary: "Revenue was {{nope}}"
sources:
  - {{name: d, type: csv, path: "{sales_csv}"}}
exports:
  output_dir: "{tmp_path / "o"}"
  filename: bad
""")
    with pytest.raises(Exception, match="unknown metric"):
        run_report(config, formats=["html"])


def test_display_cap_does_not_truncate_excel(write_config, sales_csv, tmp_path):
    """max_rows caps PDF/HTML display only; Excel keeps every row and the
    caption must state the real shown/total counts."""
    config = write_config(f"""
version: 1
report: {{title: Cap test, template: general}}
sources:
  - {{name: d, type: csv, path: "{sales_csv}"}}
sections:
  - {{title: Data, type: table, dataset: d, max_rows: 2}}
exports:
  output_dir: "{tmp_path / "o"}"
  filename: cap
  overwrite: true
""")
    outputs = run_report(config, formats=["html", "xlsx"])
    html = next(p for p in outputs if p.suffix == ".html").read_text(encoding="utf-8")
    assert "Showing 2 of 5 rows" in html
    book = openpyxl.load_workbook(next(p for p in outputs if p.suffix == ".xlsx"))
    sheet = book["Data"]
    assert sheet.max_row == 6  # header + all 5 data rows


def test_traversal_filename_contained(write_config, sales_csv, tmp_path):
    config = write_config(f"""
version: 1
report: {{title: T, template: general}}
sources:
  - {{name: d, type: csv, path: "{sales_csv}"}}
exports:
  output_dir: "{tmp_path / "o"}"
  filename: "../../escape"
  overwrite: true
""")
    outputs = run_report(config, formats=["html"])
    assert Path(outputs[0]).resolve().parent == (tmp_path / "o").resolve()
    assert outputs[0].name == "escape.html"


# ------------------------------------------------------------ fixes #15-#18


def test_excel_includes_text_sections_and_methodology(write_config, sales_csv, tmp_path):
    csv = sales_csv
    config = write_config(
        f"""
version: 1
report:
  title: Excel Parity
  template: general
  methodology: "Confidential methodology note."
sources: [{{name: d, type: csv, path: {csv}}}]
sections:
  - {{type: text, title: Notes, text: "Important context sentence."}}
  - {{type: table, title: Rows, dataset: d}}
exports: {{output_dir: {tmp_path / "out"}, filename: parity, overwrite: true}}
"""
    )
    (xlsx,) = run_report(config, formats=["xlsx"])
    book = openpyxl.load_workbook(xlsx)
    values = [str(cell.value) for row in book["Summary"].iter_rows() for cell in row if cell.value]
    assert any("Important context sentence." in v for v in values)
    assert any("Confidential methodology note." in v for v in values)


def test_excel_metric_precision_and_currency(write_config, sales_csv, tmp_path):
    csv = sales_csv
    config = write_config(
        f"""
version: 1
report: {{title: Precision, template: general, currency: "€"}}
sources: [{{name: d, type: csv, path: {csv}}}]
metrics:
  - {{name: total, dataset: d, aggregation: sum, column: revenue, format: currency, decimals: 0, label: Total}}
sections:
  - {{type: metrics, title: KPIs}}
exports: {{output_dir: {tmp_path / "out"}, filename: precision, overwrite: true}}
"""
    )
    (xlsx,) = run_report(config, formats=["xlsx"])
    book = openpyxl.load_workbook(xlsx)
    formats = [
        cell.number_format
        for row in book["Summary"].iter_rows()
        for cell in row
        if isinstance(cell.value, (int, float))
    ]
    assert any("€" in f and f.endswith("#,##0") for f in formats), formats


def test_pdf_tall_chart_and_long_cells(write_config, tmp_path):
    csv = tmp_path / "long.csv"
    csv.write_text(
        "x,y,notes\n" + "".join(f"{i},{i * 2}," + "z" * 5000 + f"{i}\n" for i in range(5))
    )
    config = write_config(
        f"""
version: 1
report: {{title: Tall, template: general}}
sources: [{{name: d, type: csv, path: {csv}}}]
sections:
  - type: chart
    title: Very tall chart
    dataset: d
    chart: {{type: line, x: x, y: y, width: 4, height: 14}}
  - {{type: table, title: Wide text, dataset: d}}
exports: {{output_dir: {tmp_path / "out"}, filename: tall, overwrite: true}}
"""
    )
    (pdf,) = run_report(config, formats=["pdf"])
    assert pdf.read_bytes().startswith(b"%PDF")


def test_partial_outputs_survive_later_format_failure(write_config, sales_csv, tmp_path):
    csv = sales_csv
    out = tmp_path / "out"
    out.mkdir()
    (out / "partial.pdf").write_text("existing")  # pdf will refuse to overwrite
    config = write_config(
        f"""
version: 1
report: {{title: Partial, template: general}}
sources: [{{name: d, type: csv, path: {csv}}}]
exports: {{output_dir: {out}, filename: partial}}
"""
    )
    from reportgen.errors import PartialExportError

    with pytest.raises(PartialExportError) as info:
        run_report(config, formats=["html", "pdf"])
    assert [p.name for p in info.value.outputs] == ["partial.html"]
    assert (out / "partial.html").exists()
    assert (out / "partial.pdf").read_text() == "existing"


def test_currency_in_metric_text(write_config, sales_csv, tmp_path):
    csv = sales_csv
    config = write_config(
        f"""
version: 1
report: {{title: Euro, template: general, currency: "€"}}
sources: [{{name: d, type: csv, path: {csv}}}]
metrics:
  - {{name: total, dataset: d, aggregation: sum, column: revenue, format: currency}}
sections: [{{type: metrics, title: KPIs}}]
exports: {{output_dir: {tmp_path / "out"}, filename: euro, overwrite: true}}
"""
    )
    (html,) = run_report(config, formats=["html"])
    text = html.read_text(encoding="utf-8")
    assert "€890.00" in text
