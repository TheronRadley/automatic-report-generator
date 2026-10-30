"""Excel exporter built on OpenPyXL.

Values keep their native types (numbers stay numbers, dates stay dates)
so readers can compute with them. String cells are passed through
``excel_safe_value`` to neutralise formula injection.
"""

from __future__ import annotations

import re
from datetime import date, datetime
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from reportgen.errors import ExportError
from reportgen.reports.model import ReportDocument
from reportgen.utils.security import excel_safe_value

_HEADER_FILL = PatternFill("solid", fgColor="2F4B7C")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=10)
_TITLE_FONT = Font(bold=True, size=14)
_LABEL_FONT = Font(color="5A5A5A", size=9)


def _metric_format(number_format: str, decimals: int, currency: str) -> str:
    """Excel number format honoring the metric's configured precision
    and the report's currency."""
    decimals = max(0, int(decimals))
    fraction = ("." + "0" * decimals) if decimals else ""
    if number_format == "currency":
        safe_currency = currency.replace('"', "")
        return f'"{safe_currency}"#,##0{fraction}'
    if number_format == "percent":
        return f"0{fraction}%"
    if number_format == "integer":
        return "#,##0"
    return f"#,##0{fraction}"


def _write_text_block(sheet, row: int, heading: str, text: str) -> int:
    """Write a heading plus wrapped text; return the next free row."""
    sheet.cell(row=row, column=1, value=heading).font = Font(bold=True)
    row += 1
    lines = [part for part in str(text).split("\n") if part.strip()] or [str(text)]
    for line in lines:
        sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        cell = sheet.cell(row=row, column=1, value=excel_safe_value(line))
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        estimated_lines = max(1, len(line) // 110 + 1)
        sheet.row_dimensions[row].height = min(180, 15 * estimated_lines)
        row += 1
    return row + 1


def _to_excel(value):
    if value is None:
        return None
    try:
        if pd.isna(value):
            return None
    except (TypeError, ValueError):
        pass
    if isinstance(value, pd.Timestamp):
        return value.to_pydatetime()
    if isinstance(value, (datetime, date)):
        return value
    if hasattr(value, "item"):  # numpy scalars -> Python scalars
        try:
            return excel_safe_value(value.item())
        except (ValueError, AttributeError):
            pass
    if isinstance(value, bool):
        return value
    return excel_safe_value(value)


def _sheet_name(title: str, used: set[str]) -> str:
    cleaned = re.sub(r"[\[\]:*?/\\]", " ", title).strip() or "Sheet"
    cleaned = cleaned[:31]
    name = cleaned
    counter = 2
    while name.lower() in used:
        suffix = f" {counter}"
        name = cleaned[: 31 - len(suffix)] + suffix
        counter += 1
    used.add(name.lower())
    return name


def _style_header_row(sheet, row: int, ncols: int) -> None:
    for col in range(1, ncols + 1):
        cell = sheet.cell(row=row, column=col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="left", vertical="center")


def _auto_widths(sheet, frame: pd.DataFrame, header_row: int) -> None:
    for index, column in enumerate(frame.columns, start=1):
        header_len = len(str(column))
        sample = frame[column].head(200).astype(str)
        content_len = int(sample.str.len().max()) if len(sample) else 0
        width = min(48, max(11, max(header_len, content_len) + 2))
        sheet.column_dimensions[get_column_letter(index)].width = width
    sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1).coordinate


def write_excel(document: ReportDocument, path: Path) -> None:
    try:
        workbook = Workbook()
        used_names: set[str] = set()

        summary = workbook.active
        summary.title = _sheet_name("Summary", used_names)
        summary["A1"] = excel_safe_value(document.title)
        summary["A1"].font = _TITLE_FONT
        row = 2
        if document.subtitle:
            summary.cell(row=row, column=1, value=excel_safe_value(document.subtitle))
            row += 1
        meta = [("Generated", document.generated_label)]
        if document.author:
            meta.append(("Author", document.author))
        if document.organization:
            meta.append(("Organization", document.organization))
        meta.append(("Template", document.template))
        for label, value in meta:
            summary.cell(row=row, column=1, value=label).font = _LABEL_FONT
            summary.cell(row=row, column=2, value=excel_safe_value(value))
            row += 1
        row += 1
        summary.cell(row=row, column=1, value="Executive summary").font = Font(bold=True)
        row += 1
        summary.merge_cells(start_row=row, start_column=1, end_row=row, end_column=4)
        cell = summary.cell(row=row, column=1, value=excel_safe_value(document.summary))
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        summary.row_dimensions[row].height = 58
        row += 2

        if document.metrics:
            summary.cell(row=row, column=1, value="Key metrics").font = Font(bold=True)
            row += 1
            header_row = row
            summary.cell(row=row, column=1, value="Metric")
            summary.cell(row=row, column=2, value="Value")
            summary.cell(row=row, column=3, value="Dataset")
            _style_header_row(summary, header_row, 3)
            row += 1
            for metric in document.metrics:
                summary.cell(row=row, column=1, value=excel_safe_value(metric.label))
                value_cell = summary.cell(row=row, column=2, value=metric.value)
                value_cell.number_format = _metric_format(
                    metric.number_format, metric.decimals, document.currency
                )
                summary.cell(row=row, column=3, value=excel_safe_value(metric.dataset))
                row += 1
            row += 1

        # Parity with HTML/PDF: text sections and methodology are part
        # of the deliverable in every format.
        for section in document.sections:
            if section.kind == "text" and (section.text or section.description):
                row = _write_text_block(
                    summary, row, section.title, section.text or section.description or ""
                )
        if document.methodology:
            row = _write_text_block(summary, row, "Methodology", document.methodology)

        summary.cell(row=row, column=1, value="Sources").font = Font(bold=True)
        row += 1
        for source in document.sources:
            summary.cell(row=row, column=1, value=excel_safe_value(source))
            row += 1
        if document.warnings:
            row += 1
            summary.cell(row=row, column=1, value="Warnings").font = Font(bold=True)
            row += 1
            for warning in document.warnings:
                summary.cell(row=row, column=1, value=excel_safe_value(warning))
                row += 1
        summary.column_dimensions["A"].width = 34
        summary.column_dimensions["B"].width = 22
        summary.column_dimensions["C"].width = 24
        summary.column_dimensions["D"].width = 18

        for section in document.sections:
            if section.kind == "table" and section.table is not None:
                sheet = workbook.create_sheet(_sheet_name(section.title, used_names))
                frame = section.table
                for col_index, column in enumerate(frame.columns, start=1):
                    sheet.cell(row=1, column=col_index, value=excel_safe_value(str(column)))
                _style_header_row(sheet, 1, len(frame.columns))
                for row_index, row_values in enumerate(frame.itertuples(index=False), start=2):
                    for col_index, value in enumerate(row_values, start=1):
                        cell = sheet.cell(row=row_index, column=col_index, value=_to_excel(value))
                        if isinstance(cell.value, float):
                            cell.number_format = "#,##0.00"
                        elif isinstance(cell.value, (datetime, date)):
                            cell.number_format = "yyyy-mm-dd"
                if len(frame.columns):
                    sheet.auto_filter.ref = (
                        f"A1:{get_column_letter(len(frame.columns))}{len(frame) + 1}"
                    )
                _auto_widths(sheet, frame, header_row=1)
            elif section.kind == "chart" and section.chart_path and section.chart_path.exists():
                from openpyxl.drawing.image import Image as ExcelImage

                sheet = workbook.create_sheet(_sheet_name(section.title, used_names))
                sheet["A1"] = excel_safe_value(section.chart_title or section.title)
                sheet["A1"].font = Font(bold=True)
                image = ExcelImage(str(section.chart_path))
                scale = min(1.0, 1000 / max(image.width, 1))
                image.width = int(image.width * scale)
                image.height = int(image.height * scale)
                sheet.add_image(image, "A3")

        workbook.save(str(path))
    except ExportError:
        raise
    except Exception as exc:
        raise ExportError(f"Excel rendering failed: {exc}") from exc
